"""Print da planilha + texto calculado: o PHD (e parecidos) sem filtrar, ocultar e copiar à mão.

Lê uma aba do Excel, separa as linhas por um valor (ex.: supervisor), mostra só as
colunas escolhidas e os dias da semana de referência, desenha a tabela numa imagem
e calcula os campos do texto ({acima}, {piores}...). Cada valor vira uma "parte":
uma imagem com a sua legenda.

As partes de uma rodada ficam gravadas numa pasta com um manifest.json. Assim, se o
PC reiniciar no meio da rodada, quem ainda não recebeu recebe exatamente a mesma
imagem e o mesmo texto de quem já recebeu.
"""

import json
import os
import re
import shutil
import tempfile
import time
from dataclasses import asdict, dataclass
from datetime import date, datetime, timedelta
from functools import lru_cache
from pathlib import Path

from . import rules
from .history import analyst_slug
from .models import Dispatch, Snapshot

WEEKDAYS = ["SEG", "TER", "QUA", "QUI", "SEX", "SÁB", "DOM"]
KEEP_DAYS = 30  # pastas de rodadas mais velhas que isso são apagadas


class SnapshotError(Exception):
    """A planilha não tem o que o envio espera (aba, coluna, linhas). A mensagem vai para a tela."""


@dataclass(frozen=True)
class Column:
    letter: str
    title: str
    day: date | None = None  # preenchido nas colunas de dia


@dataclass
class Group:
    value: str
    columns: list[Column]
    rows: list[list]  # valores alinhados com columns
    names: list[str]
    metrics: list[float | None]


@dataclass(frozen=True)
class Part:
    key: str  # estável entre rodadas: vira parte do identificador do destinatário no histórico
    label: str
    subject: str
    text: str
    image: Path


# ---------- leitura ----------

def col_index(letter: str) -> int:
    """'A' -> 0, 'AJ' -> 35."""
    letter = letter.strip().upper()
    if not re.fullmatch(r"[A-Z]{1,3}", letter):
        raise SnapshotError(f"coluna inválida: {letter!r}")
    n = 0
    for ch in letter:
        n = n * 26 + ord(ch) - 64
    return n - 1


def col_letter(index: int) -> str:
    out = ""
    index += 1
    while index:
        index, rem = divmod(index - 1, 26)
        out = chr(65 + rem) + out
    return out


def parse_columns(text: str) -> list[str]:
    """'B, D:E, AJ' -> ['B', 'D', 'E', 'AJ']. Lança SnapshotError se algo não for coluna."""
    out = []
    for part in re.split(r"[,;\s]+", text.strip()):
        if not part:
            continue
        if ":" in part:
            a, b = part.split(":", 1)
            ia, ib = col_index(a), col_index(b)
            out += [col_letter(i) for i in range(min(ia, ib), max(ia, ib) + 1)]
        else:
            out.append(col_letter(col_index(part)))
    return out


def to_number(value) -> float | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if "," in text:  # "1.234,50" escrito como texto
        text = text.replace(".", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def _as_day(value, ref: date) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    n = to_number(value)
    if n is None or n != int(n) or not 1 <= n <= 31:
        return None
    try:
        return ref.replace(day=int(n))
    except ValueError:  # dia 31 num mês de 30
        return None


def reference_day(cfg: Snapshot, today: date) -> date:
    return today if cfg.reference == "hoje" else rules.last_sales_day(today)


def read_groups(cfg: Snapshot, today: date) -> tuple[date, list[Group]]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(cfg.source, read_only=True, data_only=True)
    except FileNotFoundError:
        raise SnapshotError(f"planilha não encontrada: {cfg.source}") from None
    except Exception as exc:
        raise SnapshotError(f"não consegui abrir {Path(cfg.source).name}: {exc}") from exc
    try:
        if cfg.sheet:
            if cfg.sheet not in wb.sheetnames:
                raise SnapshotError(f"aba \"{cfg.sheet}\" não existe em {Path(cfg.source).name}")
            ws = wb[cfg.sheet]
        else:
            ws = wb.worksheets[0]
        grid = [list(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()

    def cell(row: int, letter: str):
        if row < 1 or row > len(grid):
            return None
        values, i = grid[row - 1], col_index(letter)
        return values[i] if i < len(values) else None

    ref = reference_day(cfg, today)
    fixed = [Column(c, _text(cell(cfg.header_row, c)) or c) for c in cfg.columns]
    extras = [Column(c, _text(cell(cfg.header_row, c)) or c) for c in cfg.extra_columns]
    days: list[Column] = []
    if cfg.day_columns:
        first = ref - timedelta(days=ref.weekday())
        last = min(first + timedelta(days=6), today)
        for c in parse_columns(cfg.day_columns):
            d = _as_day(cell(cfg.day_row, c), ref)
            if d and first <= d <= last and not (cfg.hide_sundays and d.weekday() == 6):
                days.append(Column(c, WEEKDAYS[d.weekday()], d))
        days.sort(key=lambda col: col.day)
    columns = fixed + days + extras
    if not columns:
        raise SnapshotError("escolha pelo menos uma coluna para o print")

    wanted = [v.strip().casefold() for v in cfg.split_values if v.strip()]
    groups: dict[str, Group] = {}
    for r in range(cfg.header_row + 1, len(grid) + 1):
        values = [cell(r, col.letter) for col in columns]
        if all(_text(cell(r, col.letter)) == "" for col in fixed or columns):
            continue
        value = _text(cell(r, cfg.split_column)) if cfg.split_column else ""
        if cfg.split_column and not value:
            continue
        if wanted and value.casefold() not in wanted:
            continue
        if cfg.metric_column:
            metric = to_number(cell(r, cfg.metric_column))
        else:
            counted = [to_number(cell(r, col.letter)) or 0.0 for col in days if col.day <= ref]
            metric = sum(counted) / len(counted) if counted else None
        name = _text(cell(r, cfg.name_column)) if cfg.name_column else _text(values[0])
        g = groups.setdefault(value.casefold(), Group(value, columns, [], [], []))
        g.rows.append(values)
        g.names.append(name)
        g.metrics.append(metric)

    ordered = list(groups.values())
    if wanted:
        ordered.sort(key=lambda g: wanted.index(g.value.casefold()))
    return ref, ordered


# ---------- texto ----------

def fmt_number(n: float | None, decimals: int = 1) -> str:
    if n is None:
        return ""
    if n == int(n):
        return str(int(n))
    return f"{n:.{decimals}f}".replace(".", ",")


def summary(group: Group, cfg: Snapshot, ref: date) -> dict[str, str]:
    scored = [(m, name) for m, name in zip(group.metrics, group.names) if m is not None]
    n = max(cfg.worst_count, 0)
    worst = sorted(scored, key=lambda x: (x[0], x[1]))[:n]
    best = sorted(scored, key=lambda x: (-x[0], x[1]))[:n]
    ref_col = next((i for i, c in enumerate(group.columns) if c.day == ref), None)
    total_ref = sum(to_number(row[ref_col]) or 0 for row in group.rows) if ref_col is not None else None
    return {
        "valor": group.value,
        "referencia": ref.strftime("%d/%m/%Y"),
        "linhas": str(len(group.rows)),
        "acima": str(sum(1 for m, _ in scored if m >= cfg.threshold)),
        "abaixo": str(sum(1 for m, _ in scored if m < cfg.threshold)),
        "limite": fmt_number(cfg.threshold),
        "qtd_piores": str(len(worst)),
        "piores": "\n".join(f"🔻 {name}" for _, name in worst),
        "melhores": "\n".join(f"🔺 {name}" for _, name in best),
        "media": fmt_number(sum(m for m, _ in scored) / len(scored)) if scored else "",
        "total_dia": fmt_number(total_ref) if total_ref is not None else "",
    }


FIELDS = ("valor", "referencia", "linhas", "acima", "abaixo", "limite", "qtd_piores", "piores",
          "melhores", "media", "total_dia")


# ---------- imagem ----------

SCALE = 2  # desenha em dobro: fica nítido no celular
NAVY, NAVY_TEXT = (11, 42, 58), (255, 255, 255)
TOTAL_BG = (89, 89, 89)
GRID, BODY_TEXT, ZEBRA = (217, 217, 217), (31, 31, 31), (247, 249, 251)
GOOD = ((198, 239, 206), (30, 107, 52))
MID = ((252, 228, 214), (197, 90, 17))
ZERO = ((253, 233, 233), (222, 150, 150))


@lru_cache(maxsize=8)
def _font(bold: bool, size: int):
    from PIL import ImageFont

    names = (["segoeuib.ttf", "arialbd.ttf", "DejaVuSans-Bold.ttf", "LiberationSans-Bold.ttf"] if bold
             else ["segoeui.ttf", "arial.ttf", "DejaVuSans.ttf", "LiberationSans-Regular.ttf"])
    for name in names:
        try:
            return ImageFont.truetype(name, size)
        except OSError:
            continue
    return ImageFont.load_default(size)


def _cell_text(value) -> str:
    if isinstance(value, float):
        return fmt_number(value, 2) if value != int(value) else str(int(value))
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y")
    return _text(value)


def day_colors(value, threshold: float) -> tuple[tuple, tuple]:
    n = to_number(value) or 0
    return GOOD if n >= threshold else MID if n > 0 else ZERO


def render_png(group: Group, cfg: Snapshot, title: str, path: Path) -> None:
    from PIL import Image, ImageDraw

    s = SCALE
    body, bold, big = _font(False, 13 * s), _font(True, 13 * s), _font(True, 17 * s)
    pad, row_h = 9 * s, 24 * s
    head_h = row_h * 2 if any(c.day for c in group.columns) else row_h
    title_h = 40 * s if title else 0
    texts = [[_cell_text(v) for v in row] for row in group.rows]

    totals = []
    for i, c in enumerate(group.columns):
        totals.append(fmt_number(sum(to_number(row[i]) or 0 for row in group.rows)) if c.day else "")
    widths = []
    for i, c in enumerate(group.columns):
        cand = [bold.getlength(c.title), bold.getlength(totals[i])] + [body.getlength(t[i]) for t in texts]
        if c.day:
            cand.append(bold.getlength(f"{c.day.day:02d}"))
        widths.append(int(max(cand + [12 * s])) + 2 * pad)
    width = sum(widths)
    if title:
        width = max(width, int(big.getlength(title)) + 4 * pad)
    has_total = any(totals)
    height = title_h + head_h + (row_h if has_total else 0) + row_h * len(texts) + s

    img = Image.new("RGB", (width, height), (255, 255, 255))
    draw = ImageDraw.Draw(img)

    def put(x0, y0, w, h, text, font, fill, align="center"):
        tw = font.getlength(text)
        x = x0 + pad if align == "left" else x0 + (w - tw) / 2
        draw.text((x, y0 + h / 2), text, font=font, fill=fill, anchor="lm")

    y = 0
    if title:
        draw.rectangle((0, 0, width, title_h), fill=NAVY)
        put(0, 0, width, title_h, title, big, NAVY_TEXT, align="left")
        y = title_h

    draw.rectangle((0, y, width, y + head_h), fill=NAVY)
    x = 0
    for c, w in zip(group.columns, widths):
        if c.day:
            put(x, y, w, row_h, c.title, bold, NAVY_TEXT)
            put(x, y + row_h, w, row_h, f"{c.day.day:02d}", bold, NAVY_TEXT)
        else:
            put(x, y, w, head_h, c.title, bold, NAVY_TEXT)
        x += w
    y += head_h

    if has_total:
        draw.rectangle((0, y, width, y + row_h), fill=TOTAL_BG)
        x = 0
        for t, w in zip(totals, widths):
            if t:
                put(x, y, w, row_h, t, bold, NAVY_TEXT)
            x += w
        y += row_h

    numeric = [bool(c.day) or all(to_number(row[i]) is not None or row[i] is None for row in group.rows)
               for i, c in enumerate(group.columns)]
    for n, (row, line) in enumerate(zip(group.rows, texts)):
        if n % 2:
            draw.rectangle((0, y, width, y + row_h), fill=ZEBRA)
        x = 0
        for i, (c, w) in enumerate(zip(group.columns, widths)):
            fg = BODY_TEXT
            if c.day:
                bg, fg = day_colors(row[i], cfg.threshold)
                draw.rectangle((x + 1, y + 1, x + w - 1, y + row_h - 1), fill=bg)
                text = line[i] or "0"
            else:
                text = line[i]
            put(x, y, w, row_h, text, bold if c.day else body, fg, align="center" if numeric[i] else "left")
            x += w
        draw.line((0, y + row_h, width, y + row_h), fill=GRID, width=1)
        y += row_h

    x = 0
    for w in widths[:-1]:
        x += w
        draw.line((x, title_h, x, height), fill=GRID, width=1)
    path.parent.mkdir(parents=True, exist_ok=True)
    img.save(path, "PNG", optimize=True)


# ---------- partes de uma rodada ----------

def _slug(text: str) -> str:
    return analyst_slug(re.sub(r"[^\w\s-]", " ", text))[:48] or "tudo"


def build_parts(d: Dispatch, now: datetime, out_dir: Path) -> list[Part]:
    cfg = d.snapshot
    ref, groups = read_groups(cfg, now.date())
    if not groups:
        what = f" para {', '.join(cfg.split_values)}" if cfg.split_values else ""
        raise SnapshotError(f"nenhuma linha encontrada na planilha{what}")
    out_dir.mkdir(parents=True, exist_ok=True)
    parts, seen = [], set()
    for i, g in enumerate(groups, start=1):
        key = _slug(g.value)
        if key in seen:
            key = f"{key}-{i}"
        seen.add(key)
        fields = summary(g, cfg, ref)
        image = out_dir / f"{i:02d}-{key}.png"
        render_png(g, cfg, rules.render(cfg.title, now, [], fields).strip(), image)
        parts.append(Part(key, g.value, rules.render(d.subject, now, [cfg.source], fields),
                          rules.render(d.message, now, [cfg.source], fields), image))
    return parts


def _write_manifest(folder: Path, parts: list[Part]) -> None:
    data = [{**asdict(p), "image": p.image.name} for p in parts]
    fd, tmp = tempfile.mkstemp(dir=folder, suffix=".tmp")
    with os.fdopen(fd, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    os.replace(tmp, folder / "manifest.json")


def _read_manifest(folder: Path) -> list[Part] | None:
    try:
        data = json.loads((folder / "manifest.json").read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None
    return [Part(p["key"], p["label"], p["subject"], p["text"], folder / p["image"]) for p in data]


class SnapshotRuns:
    """As partes geradas em cada rodada: geradas uma vez, reaproveitadas até o fim da rodada."""

    def __init__(self, base_dir: Path) -> None:
        self.base_dir = Path(base_dir)

    def folder(self, dispatch_id: str, run: str) -> Path:
        return self.base_dir / dispatch_id / re.sub(r"[^\w-]+", "-", run).strip("-")

    def load(self, dispatch_id: str, run: str) -> list[Part] | None:
        return _read_manifest(self.folder(dispatch_id, run))

    def build(self, d: Dispatch, run: str, now: datetime) -> list[Part]:
        folder = self.folder(d.id, run)
        if folder.exists():
            shutil.rmtree(folder, ignore_errors=True)
        parts = build_parts(d, now, folder)
        _write_manifest(folder, parts)  # por último: manifest presente = rodada completa
        self.cleanup()
        return parts

    def preview(self, d: Dispatch, now: datetime) -> list[Part]:
        folder = self.base_dir / "previa" / d.id
        shutil.rmtree(folder, ignore_errors=True)
        return build_parts(d, now, folder)

    def cleanup(self, keep_days: int = KEEP_DAYS) -> None:
        limit = time.time() - keep_days * 86400
        for dispatch_dir in self.base_dir.glob("*"):
            if not dispatch_dir.is_dir() or dispatch_dir.name == "previa":
                continue
            for run_dir in dispatch_dir.glob("*"):
                try:
                    if run_dir.is_dir() and run_dir.stat().st_mtime < limit:
                        shutil.rmtree(run_dir, ignore_errors=True)
                except OSError:
                    pass
