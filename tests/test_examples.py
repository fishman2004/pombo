import csv
import os

import pytest

from tools.make_examples import COLUMNS, make_examples

ROTINA_8081 = [
    "ID RCA", "NOME RCA", "NOME SUPERVISOR", "NUMERO PEDIDO", "POSICAO PED", "DATA PEDIDO", "ID CLIENTE",
    "CLIENTE", "ID PRODUTO", "NOME PRODUTO", "ID FORNECEDOR", "NOME FORNECEDOR", "QTDE CX", "QTDE UND",
    "VALOR", "QTDE CX VENDIDA",
]


def read_csv(path):
    with path.open(encoding="utf-8-sig", newline="") as f:
        return list(csv.reader(f, delimiter=";"))


def test_columns_match_rotina_8081(tmp_path):
    assert list(COLUMNS) == ROTINA_8081
    out = make_examples(tmp_path, rows=50)
    rows = read_csv(next(p for p in out if p.suffix == ".csv"))
    assert rows[0] == ROTINA_8081 and len(rows) == 51


def test_brazilian_formats(tmp_path):
    rows = read_csv(next(p for p in make_examples(tmp_path, rows=20) if p.suffix == ".csv"))
    data_pedido, valor = rows[1][5], rows[1][14]
    assert data_pedido[2] == "/" and data_pedido[5] == "/"
    assert "," in valor and "." not in valor


def test_same_seed_same_output(tmp_path):
    a = make_examples(tmp_path / "a", rows=30, seed=7)
    b = make_examples(tmp_path / "b", rows=30, seed=7)
    csv_a = next(p for p in a if p.suffix == ".csv").read_bytes()
    csv_b = next(p for p in b if p.suffix == ".csv").read_bytes()
    assert csv_a == csv_b


def test_spreadsheets_created(tmp_path):
    from openpyxl import load_workbook

    out = make_examples(tmp_path, rows=40)
    xlsx = [p for p in out if p.suffix == ".xlsx"]
    assert {p.name for p in xlsx} == {"Vendas.xlsx", "Plano_Excelencia.xlsx"}
    ws = load_workbook(tmp_path / "Vendas.xlsx").active
    assert ws["A1"].value == "NOME RCA" and ws.max_row > 1


@pytest.mark.skipif(not os.environ.get("EXEMPLOS_PROIBIDOS"), reason="lista de nomes reais só existe no PC do analista")
def test_no_real_names(tmp_path):
    banned = [n.strip().lower() for n in os.environ["EXEMPLOS_PROIBIDOS"].split(";") if n.strip()]
    text = "".join(p.read_bytes().decode("utf-8", "ignore") for p in make_examples(tmp_path) if p.suffix == ".csv").lower()
    assert not [n for n in banned if n in text]
