"""Gera dados de exemplo FICTÍCIOS no formato da rotina 8081 do Winthor.

Todos os nomes (vendedores, supervisores, clientes, produtos, fornecedores) são
inventados. Os arquivos reais da empresa nunca entram no repositório.

    .venv\\Scripts\\python -m tools.make_examples
"""

import csv
import random
from collections import defaultdict
from datetime import date, timedelta
from pathlib import Path

COLUMNS = (
    "ID RCA", "NOME RCA", "NOME SUPERVISOR", "NUMERO PEDIDO", "POSICAO PED", "DATA PEDIDO", "ID CLIENTE",
    "CLIENTE", "ID PRODUTO", "NOME PRODUTO", "ID FORNECEDOR", "NOME FORNECEDOR", "QTDE CX", "QTDE UND",
    "VALOR", "QTDE CX VENDIDA",
)

SUPERVISORES = ["S-NORTE (MARINA COUTO)", "S-SUL (RAFAEL DUARTE)", "S-CENTRO (LUANA PRADO)"]
RCAS = [
    (101, "R-01- ANDRE FALCAO NERY", 0), (102, "R-02- BEATRIZ LOBO TAVARES", 0),
    (103, "R-03- CAIO MENDONCA REIS", 1), (104, "R-04- DENISE AMARAL FONTES", 1),
    (105, "R-05- EDUARDO PIRES SALES", 2), (106, "R-06- FLAVIA ROCHA MOTA", 2),
]
CLIENTES = [
    "MERCADINHO BOM PRECO", "SUPERMERCADO ESTRELA DO VALE", "MERCEARIA DONA ROSA", "ATACADO SOL NASCENTE",
    "MINIMERCADO DOIS IRMAOS", "PADARIA PAO DOURADO", "CONVENIENCIA RAPIDINHO", "EMPORIO SAO JORGE",
    "MERCADO POPULAR DO BAIRRO", "SUPER COMPRE BEM", "DISTRIBUIDORA PRAIA AZUL", "MERCADINHO FAMILIA",
]
FORNECEDORES = [(9001, "DOCES AURORA"), (9002, "CONFEITOS BRISA"), (9003, "CHOCOLATES SERRA FRIA"), (9004, "SALGADINHOS CROC")]
PRODUTOS = [
    (5001, "BALA DE GOMA FRUTAS 500G", 9001, 24, 7.90), (5002, "PIRULITO SORTIDO 50UN", 9001, 12, 11.50),
    (5003, "CHICLETE HORTELA 100UN", 9002, 18, 9.40), (5004, "BALA DE COCO 700G", 9002, 20, 8.75),
    (5005, "CHOCOLATE AO LEITE 90G", 9003, 30, 5.60), (5006, "BOMBOM SORTIDO 250G", 9003, 16, 14.20),
    (5007, "SALGADINHO QUEIJO 45G", 9004, 40, 2.80), (5008, "AMENDOIM CROCANTE 150G", 9004, 24, 4.30),
]


def _money(value: float) -> str:
    return f"{value:.2f}".replace(".", ",")


def make_examples(out_dir: Path, rows: int = 500, seed: int = 42) -> list[Path]:
    rnd = random.Random(seed)
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    clientes = {c: 250000 + i * 137 for i, c in enumerate(CLIENTES)}
    fornecedores = dict(FORNECEDORES)
    start = date(2026, 10, 1)

    lines = []
    pedido = 429000000
    while len(lines) < rows:
        pedido += rnd.randint(1, 9)
        rca_id, rca_nome, sup = rnd.choice(RCAS)
        cliente = rnd.choice(CLIENTES)
        dia = start + timedelta(days=rnd.randint(0, 27))
        for _ in range(rnd.randint(1, 4)):
            pid, pnome, fid, und_cx, preco = rnd.choice(PRODUTOS)
            cx = rnd.randint(1, 12)
            lines.append([
                rca_id, rca_nome, SUPERVISORES[sup], pedido, rnd.choice("FFFFLM"), dia.strftime("%d/%m/%Y"),
                clientes[cliente], cliente, pid, pnome, fid, fornecedores[fid], cx, cx * und_cx,
                _money(cx * und_cx * preco), cx,
            ])
    lines = lines[:rows]

    csv_path = out_dir / "(DADOS) - 8081 - Exemplo.csv"
    with csv_path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f, delimiter=";")
        writer.writerow(COLUMNS)
        writer.writerows(lines)

    return [csv_path, *_summaries(out_dir, lines)]


def _summaries(out_dir: Path, lines: list[list]) -> list[Path]:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    def value(line) -> float:
        return float(line[14].replace(",", "."))

    by_rca: dict[str, list[float]] = defaultdict(lambda: [0, 0.0])
    pedidos: dict[str, set] = defaultdict(set)
    by_forn: dict[str, list[float]] = defaultdict(lambda: [0, 0.0])
    for line in lines:
        pedidos[line[1]].add(line[3])
        by_rca[line[1]][1] += value(line)
        by_forn[line[11]][0] += line[12]
        by_forn[line[11]][1] += value(line)
    for rca in by_rca:
        by_rca[rca][0] = len(pedidos[rca])

    def save(name: str, header: list[str], data: list[list]) -> Path:
        wb = Workbook()
        ws = wb.active
        ws.title = "Resumo"
        ws.append(header)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="C00000")
        for row in data:
            ws.append(row)
        for col, width in zip("ABC", (34, 14, 16)):
            ws.column_dimensions[col].width = width
        path = out_dir / name
        wb.save(path)
        return path

    vendas = sorted(([k, int(v[0]), round(v[1], 2)] for k, v in by_rca.items()), key=lambda r: -r[2])
    plano = sorted(([k, int(v[0]), round(v[1], 2)] for k, v in by_forn.items()), key=lambda r: -r[2])
    return [
        save("Vendas.xlsx", ["NOME RCA", "PEDIDOS", "VALOR TOTAL"], vendas),
        save("Plano_Excelencia.xlsx", ["NOME FORNECEDOR", "QTDE CX", "VALOR"], plano),
    ]


if __name__ == "__main__":
    for p in make_examples(Path(__file__).resolve().parent.parent / "exemplos"):
        print("gerado:", p)
