"""Gera um PHD de exemplo FICTÍCIO, no mesmo desenho da planilha de verdade.

Serve para testar o envio com print da planilha sem usar dados da empresa:
título no topo, linha 13 com o dia da semana, linha 14 com o dia do mês,
linha 15 com os nomes das colunas e os totais, e um vendedor por linha a partir
da 16 (positivações de cada dia de F a AJ e a média em AK).

    .venv\\Scripts\\python -m tools.make_phd
"""

import random
from calendar import monthrange
from datetime import date
from pathlib import Path

SEMANA = ["SEG", "TER", "QUA", "QUI", "SEX", "SÁB", "DOM"]
# (área, supervisor) e quantos vendedores cada um tem; todos os nomes são inventados
EQUIPES = [
    ("LITORAL", "S-LT (MARINA COUTO)", 8),
    ("AGRESTE", "S-AG (RAFAEL DUARTE)", 9),
    ("SERTÃO", "S-ST (LUANA PRADO)", 12),
]
NOMES = [
    "ANDRE FALCAO NERY", "BEATRIZ LOBO TAVARES", "CAIO MENDONCA REIS", "DENISE AMARAL FONTES",
    "EDUARDO PIRES SALES", "FLAVIA ROCHA MOTA", "GUSTAVO LEME BRANDAO", "HELENA CASTRO VIANA",
    "IGOR TEIXEIRA LOPES", "JULIANA MORAES PINTO", "KLEBER AZEVEDO CRUZ", "LARISSA FARIAS MELO",
    "MARCOS VIEIRA PRADO", "NATALIA CORREIA DIAS", "OTAVIO RAMOS FREIRE", "PAULA BRITO CAMPOS",
    "RENATO SOARES LIMA", "SABRINA NOGUEIRA LUZ", "TIAGO MACEDO SOUZA", "URSULA BARROS REGO",
    "VITOR CARDOSO NUNES", "WANDA PEIXOTO REIS", "YURI MONTEIRO SALES", "ZULEIDE ARAUJO MAIA",
    "ALINE COSTA BARRETO", "BRUNO MATOS QUEIROZ", "CAMILA SENA ROCHA", "DIEGO FONSECA LIRA",
    "ELAINE PORTO MOURA",
]


def make_phd(path: Path, month: date = date(2026, 10, 1), until: date = date(2026, 10, 22), seed: int = 7) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    rnd = random.Random(seed)
    wb = Workbook()
    ws = wb.active
    ws.title = "PHD"
    navy = PatternFill("solid", fgColor="0B2A3A")
    white = Font(bold=True, color="FFFFFF")
    ws["B2"] = "PHD - POSITIVAÇÃO REPRESENTANTE DIA"
    ws["B2"].font = Font(bold=True, size=20, color="FFFFFF")
    ws["B2"].fill = navy

    days = monthrange(month.year, month.month)[1]
    first_day_col = 6  # F
    for i, col in enumerate(("ÁREA", "ID", "SUPERVISOR", "AREA"), start=2):
        ws.cell(15, i, col)
    rows = []
    n = 0
    for area, sup, qtd in EQUIPES:
        prefix = sup.split(" ")[0]
        for k in range(qtd):
            name = NOMES[n % len(NOMES)]
            n += 1
            strength = rnd.uniform(1.5, 9)
            values = []
            for d in range(1, days + 1):
                day = month.replace(day=d)
                if day > until or day.weekday() == 6:
                    values.append(0)
                elif day == until:  # o dia de hoje ainda está no começo
                    values.append(max(0, int(rnd.gauss(strength / 4, 1))))
                elif day.weekday() == 5:
                    values.append(max(0, int(rnd.gauss(strength / 3, 1))))
                else:
                    values.append(max(0, int(rnd.gauss(strength, 2))))
            rows.append([area, 100 + n, f"{prefix}-{k + 1:02d}- {name}", sup, values])

    for d in range(1, days + 1):
        c = first_day_col + d - 1
        ws.cell(13, c, SEMANA[month.replace(day=d).weekday()])
        ws.cell(14, c, d)
        ws.cell(15, c, sum(r[4][d - 1] for r in rows))  # na planilha real é um SUBTOTAL
    media_col = first_day_col + days
    ws.cell(15, media_col, "MÉDIA")
    for r, (area, rid, name, sup, values) in enumerate(rows, start=16):
        ws.cell(r, 2, area)
        ws.cell(r, 3, rid)
        ws.cell(r, 4, name)
        ws.cell(r, 5, sup)
        for d, v in enumerate(values, start=1):
            ws.cell(r, first_day_col + d - 1, v)
        worked = [v for d, v in enumerate(values, start=1)
                  if month.replace(day=d) <= until and month.replace(day=d).weekday() != 6]
        ws.cell(r, media_col, round(sum(worked) / len(worked), 2) if worked else 0)

    for row in ws.iter_rows(min_row=13, max_row=15, min_col=2, max_col=media_col):
        for cell in row:
            cell.fill, cell.font, cell.alignment = navy, white, Alignment(horizontal="center")
    for col, width in (("B", 16), ("C", 6), ("D", 36), ("E", 26)):
        ws.column_dimensions[col].width = width
    for c in range(first_day_col, media_col + 1):
        ws.column_dimensions[get_column_letter(c)].width = 5
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


if __name__ == "__main__":
    print("gerado:", make_phd(Path(__file__).resolve().parent.parent / "exemplos" / "PHD_Exemplo.xlsx"))
